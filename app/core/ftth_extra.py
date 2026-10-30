# -*- coding: utf-8 -*-
"""Extras for the FTTH updater: validation, optical loss budget, change report, splitter link table."""
import heapq, math, statistics
import numpy as np
from ftth_update import L, num, ceil
from ftth_std import STD, LAY

PORTS = lambda: STD['splitter']['ports']
MINSP = lambda: STD['splitter']['min_spare_ports']

class _Loss(dict):
    """loss parameters read live from STD['loss'] (keys kept for the old code)"""
    _map = dict(fibre='fibre_db_per_km', split='splitter_db', nsplit='splitter_count', splice='splice_db', conn='connector_db', nconn='connector_count', margin='margin_db', fixed_splices='fixed_splices', factor='path_factor', limit='class_c_plus_db')
    def __getitem__(self, k): return STD['loss'][self._map[k]]
LOSS = _Loss()

def _key(p): return (round(p[0]), round(p[1]))

def cable_graph(layers):
    adj = {}; segs = []
    for nm in (LAY['main_cables'], LAY['branch_cables']):
        for r in L(layers, nm):
            g = r['geometry']
            if g is None: continue
            for ln in (g.geoms if g.geom_type.startswith('Multi') else [g]):
                c = list(ln.coords); a, b = _key(c[0]), _key(c[-1]); w = ln.length
                adj.setdefault(a, []).append((b, w)); adj.setdefault(b, []).append((a, w))
    return adj

def _nearest(nodes, A, pt, tol):
    if not len(A): return None
    d = np.hypot(A[:, 0] - pt.x, A[:, 1] - pt.y); k = int(np.argmin(d))
    return nodes[k] if d[k] <= tol else None

def loss_budget(layers, tol=None):
    """per secondary splitter: shortest cable path OLT->splitter (x1.2), closures on path, loss dB.  Returns (rows, unreachable_ids)"""
    tol = tol if tol is not None else STD['snap']['splitter_to_cable_m']
    adj = cable_graph(layers); S = L(layers, LAY['secondary_splitters']); co = L(layers, LAY['center'])
    if not adj or not co or co[0]['geometry'] is None: return [], [s.get('Secondary_ID') for s in S]
    nodes = list(adj); A = np.array(nodes, float)
    o = _nearest(nodes, A, co[0]['geometry'], STD['snap']['olt_m'])
    if o is None: return [], [s.get('Secondary_ID') for s in S]
    d = {o: 0.0}; pr = {}; h = [(0.0, o)]
    while h:
        x, u = heapq.heappop(h)
        if x > d[u]: continue
        for v, w in adj[u]:
            if x + w < d.get(v, 1e18): d[v] = x + w; pr[v] = u; heapq.heappush(h, (x + w, v))
    clos = []
    for n in (LAY['closure_main'], LAY['closure_branch'], LAY['closure_merge']):
        clos += [(r['geometry'].x, r['geometry'].y) for r in L(layers, n) if r['geometry'] is not None]
    cn = {}
    for (x, y) in clos:
        k = _nearest(nodes, A, type('P', (), {'x': x, 'y': y}), STD['snap']['closure_m'])
        if k is not None: cn[k] = cn.get(k, 0) + 1
    Lc = LOSS; rows = []; bad = []
    for s in S:
        g = s['geometry']; sid = s.get('Secondary_ID')
        n = _nearest(nodes, A, g, tol) if g is not None else None
        if n is None or n not in d: bad.append(sid); continue
        path = [n]
        while path[-1] != o: path.append(pr[path[-1]])
        nc = sum(1 for p in path[1:-1] if p in cn) if len(path) > 2 else 0
        km = d[n] * Lc['factor'] / 1000.0
        loss = Lc['fibre'] * km + Lc['split'] * Lc['nsplit'] + Lc['conn'] * Lc['nconn'] + Lc['splice'] * (Lc['fixed_splices'] + nc)
        rows.append(dict(id=sid, prim=s.get('Primary_ID') or s.get('Group_ID') or '-', m=round(d[n] * Lc['factor']), closures=nc, loss=round(loss, 2), withm=round(loss + Lc['margin'], 2)))
    return rows, bad

def loss_summary(rows):
    if not rows: return None
    v = [r['loss'] for r in rows]
    return dict(min=round(min(v), 2), avg=round(statistics.mean(v), 2), max=round(max(v), 2), maxm=round(max(v) + LOSS['margin'], 2), n=len(v), over=sum(1 for r in rows if r['withm'] > LOSS['limit']), over28=sum(1 for r in rows if r['withm'] > STD['loss']['class_b_plus_db']),
                longest=max(r['m'] for r in rows))

# ------------------------------------------------------------------ validation
def pkey(r): return r.get('Primary_ID') or r.get('Group_ID') or ''

def validate(layers):
    """returns list of (level, message); level 'خطأ' stops generation, 'تنبيه' only warns"""
    out = []; P = L(layers, LAY['primary_splitters']); S = L(layers, LAY['secondary_splitters'])
    M = L(layers, LAY['main_cables']); B = L(layers, LAY['branch_cables'])
    # duplicate ids
    for nm, key in ((LAY['primary_splitters'], 'Primary_ID'), (LAY['secondary_splitters'], 'Secondary_ID'), (LAY['main_cables'], 'Cable_ID'), (LAY['branch_cables'], 'Cable_ID'), (LAY['closure_main'], 'Closure_ID'), (LAY['closure_branch'], 'Closure_ID')):
        ids = [r.get(key) for r in L(layers, nm) if r.get(key)]
        d = sorted({i for i in ids if ids.count(i) > 1})
        if d: out.append(('خطأ', 'معرّفات مكررة في طبقة %s: %s' % (nm.replace('_', ' '), '، '.join(map(str, d[:8])) + (' …' if len(d) > 8 else ''))))
    # secondary splitters vs cables
    if B:
        bn = [r['geometry'] for r in B if r['geometry'] is not None]
        from shapely.ops import unary_union
        ub = unary_union(bn) if bn else None
        P0 = L(layers, LAY['primary_splitters'])  # سبلتر فرعي ملاصق لرئيسي لا يحتاج كبلاً
        bad = [s.get('Secondary_ID') for s in S if s['geometry'] is not None and ub is not None and s['geometry'].distance(ub) > STD['snap']['validate_splitter_to_cable_m'] and not any(s['geometry'].distance(q['geometry']) <= STD['snap']['validate_splitter_to_cable_m'] for q in P0 if q['geometry'] is not None)]
        if bad: out.append(('خطأ', 'سبلترات فرعية بلا كبل قريب (أبعد من 5 م): %s' % '، '.join(map(str, bad[:10])) + (' …' if len(bad) > 10 else '')))
    # splitter ports
    lowS = []; lowP = []
    for s in S:
        used = s.get('Used_Ports', s.get('Buildings')); sp = s.get('Spare_Ports')
        if used is not None and num(used) > PORTS(): out.append(('خطأ', 'السبلتر الفرعي %s يخدم %d مبنى (> %d منافذ)' % (s.get('Secondary_ID'), num(used), PORTS())))
        if used is not None and sp is not None and num(sp) < MINSP() and num(used) + num(sp) <= PORTS(): lowS.append(str(s.get('Secondary_ID')))
    for p in P:
        kids = [s for s in S if pkey(s) == p.get('Primary_ID')]
        if len(kids) > PORTS(): out.append(('خطأ', 'السبلتر الرئيسي %s مرتبط بـ %d سبلتر فرعي (> %d منافذ)' % (p.get('Primary_ID'), len(kids), PORTS())))
        elif len(kids) == 0: out.append(('تنبيه', 'السبلتر الرئيسي %s بلا سبلترات فرعية' % p.get('Primary_ID')))
        elif PORTS() - len(kids) < MINSP(): lowP.append(str(p.get('Primary_ID')))
    if lowP: out.append(('تنبيه', '%d سبلتر رئيسي بمنفذ احتياطي أقل من %d: %s' % (len(lowP), MINSP(), '، '.join(lowP[:10]) + (' …' if len(lowP) > 10 else ''))))
    if lowS: out.append(('تنبيه', '%d سبلتر فرعي بمنفذ احتياطي أقل من %d: %s' % (len(lowS), MINSP(), '، '.join(lowS[:10]) + (' …' if len(lowS) > 10 else ''))))
    for s in S:
        if s.get('Buildings') == 0 or s.get('Used_Ports') == 0: out.append(('تنبيه', 'السبلتر الفرعي %s بلا مباني مخدومة' % s.get('Secondary_ID')))
    # cable fibres vs used
    for r in M + B:
        c, u = num(r.get('Cores')), r.get('Used')
        if u not in (None, '') and num(u) > c: out.append(('خطأ', 'الكبل %s: الألياف المستعملة %d أكبر من سعته %d' % (r.get('Cable_ID'), num(u), c)))
    # buildings without splitter
    bl = L(layers, LAY['buildings'])
    if bl:
        no = [b.get('Building_ID') for b in bl if not b.get('Secondary_ID')]
        if no: out.append(('تنبيه', '%d مبنى بلا سبلتر فرعي (مثال: %s)' % (len(no), '، '.join(map(str, no[:5])))))
    try:
        import ftth_poles; out += ftth_poles.validate(layers)
    except Exception as e: out.append(('تنبيه', 'تعذّر فحص الأعمدة: %s' % e))
    return out

# ------------------------------------------------------------------ loss summary text
def loss_xlsx(rows, bad, out_path, town):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ميزانية الفقد'; ws.sheet_view.rightToLeft = True
    sm = loss_summary(rows) or {}
    ws.append(['ميزانية الفقد الضوئي المحسوبة من الكبلات الفعلية - %s' % town]); ws['A1'].font = Font(bold=True, size=13)
    ws.append(['الطريقة: أقصر مسار كبلات من المقسم OLT إلى كل سبلتر فرعي × 1.20، وكل علبة وصل على المسار = لحام واحد.'])
    ws.append([])
    for a, b in (('توهين الليف dB/km', LOSS['fibre']), ('فقد سبلتر 1:8 (×2) dB', LOSS['split']), ('فقد لحام dB', LOSS['splice']), ('فقد موصل (×3) dB', LOSS['conn']), ('لحامات ثابتة لكل مسار', LOSS['fixed_splices']), ('هامش الأمان dB', LOSS['margin']), ('الحد الأقصى المسموح مع الهامش dB (Class C+)', LOSS['limit'])): ws.append([a, b])
    ws.append([])
    for a, b in (('أقل فقد dB', sm.get('min')), ('متوسط الفقد dB', sm.get('avg')), ('أقصى فقد dB', sm.get('max')), ('أقصى فقد مع الهامش dB', sm.get('maxm')), ('عدد السبلترات الفرعية', sm.get('n')), ('عدد ما يتجاوز 28 dB مع الهامش (Class B+)', sm.get('over28')), ('عدد ما يتجاوز 32 dB مع الهامش (Class C+)', sm.get('over')), ('أطول مسار (م)', sm.get('longest')), ('سبلترات غير موصولة بالكبلات', len(bad))): ws.append([a, b])
    ws2 = wb.create_sheet('تفصيل'); ws2.sheet_view.rightToLeft = True
    ws2.append(['السبلتر الفرعي', 'السبلتر الرئيسي', 'طول المسار (م)', 'علب على المسار', 'الفقد (dB)', 'مع الهامش (dB)'])
    for c in ws2[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='14284B'); c.alignment = Alignment(horizontal='center')
    for r in sorted(rows, key=lambda r: -r['loss']): ws2.append([r['id'], r['prim'], r['m'], r['closures'], r['loss'], r['withm']])
    for col, w in zip('ABCDEF', (34, 16, 16, 16, 12, 14)): ws2.column_dimensions[col].width = w
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 16
    wb.save(out_path)

# ------------------------------------------------------------------ splitter link table
def link_table(layers, out_path):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from collections import defaultdict
    P = L(layers, LAY['primary_splitters']); S = L(layers, LAY['secondary_splitters']); B = L(layers, LAY['branch_cables'])
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ربط الرئيسي بالفرعي'; ws.sheet_view.rightToLeft = True
    hf = PatternFill('solid', fgColor='1F3864'); th = Side(style='thin', color='999999'); bd = Border(left=th, right=th, top=th, bottom=th)
    def head(w, H):
        w.append(H)
        for c in w[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = hf; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = bd
        w.row_dimensions[1].height = 32
    head(ws, ['السبلتر الرئيسي', 'عدد السبلترات الفرعية', 'المنافذ المستخدمة (من 8)', 'المنافذ الاحتياطية', 'السبلترات الفرعية المرتبطة', 'عدد المباني المخدومة', 'طول الكوابل الفرعية (م)', 'ألياف الكبل الفرعي'])
    bl = defaultdict(float); bc = defaultdict(int)
    for r in B:
        ps = [x for x in str(pkey(r)).split('+') if x]
        for q in ps: bl[q] += num(r.get('Length_m')) / max(1, len(ps)); bc[q] = max(bc[q], int(num(r.get('Cores'))))
    kids = defaultdict(list)
    for s in S: kids[pkey(s)].append(s)
    bld = lambda s: int(num(s.get('Used_Ports', s.get('Buildings'))))
    tot = [0, 0, 0.0]
    for p in sorted(P, key=lambda r: r['Primary_ID']):
        pid = p['Primary_ID']; ks = sorted(kids.get(pid, []), key=lambda r: r['Secondary_ID']); subs = sum(bld(s) for s in ks)
        ws.append([pid, len(ks), len(ks), PORTS() - len(ks), '، '.join(s['Secondary_ID'] for s in ks), subs, round(bl.get(pid, 0.0), 1), bc.get(pid, 0)])
        tot[0] += len(ks); tot[1] += subs; tot[2] += bl.get(pid, 0.0)
    ws.append(['المجموع', tot[0], '', '', '', tot[1], round(tot[2], 1), ''])
    for row in ws.iter_rows(min_row=2):
        for c in row: c.border = bd; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        row[4].alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)
    for c in ws[ws.max_row]: c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='D9E1F2')
    for col, w in zip('ABCDEFGH', (14, 14, 16, 14, 70, 14, 16, 14)): ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'
    w2 = wb.create_sheet('تفصيل السبلترات الفرعية'); w2.sheet_view.rightToLeft = True
    head(w2, ['السبلتر الفرعي', 'السبلتر الرئيسي', 'المباني المخدومة', 'المنافذ الاحتياطية', 'البعد المباشر عن الرئيسي (م)', 'Lat', 'Lon'])
    pp = {p['Primary_ID']: p['geometry'] for p in P}
    try:
        from ftth_kml import to_lonlat
    except Exception: to_lonlat = None
    for s in sorted(S, key=lambda r: (pkey(r), r['Secondary_ID'])):
        g = s['geometry']; d = round(g.distance(pp[pkey(s)])) if pkey(s) in pp and g is not None and pp[pkey(s)] is not None else ''
        lo, la = to_lonlat(g.x, g.y) if (to_lonlat and g is not None) else ('', '')
        w2.append([s['Secondary_ID'], pkey(s), bld(s), (int(num(s['Spare_Ports'])) if s.get('Spare_Ports') not in (None, '') else max(0, PORTS() - bld(s))), d, round(la, 6) if la != '' else '', round(lo, 6) if lo != '' else ''])
    for row in w2.iter_rows(min_row=2):
        for c in row: c.border = bd; c.alignment = Alignment(horizontal='center')
    for col in 'ABCDEFG': w2.column_dimensions[col].width = 16
    w2.freeze_panes = 'A2'
    wb.save(out_path)
