# -*- coding: utf-8 -*-
"""FTTH BOQ updater: reads the (edited) GeoPackage and rewrites the BOQ workbooks.
Method: new_qty = template_qty + f(edited_gpkg) - f(baseline_gpkg)  -> an unedited gpkg reproduces the delivered BOQ exactly."""
import sys, os, json, sqlite3, struct, math, re, datetime
import numpy as np
import openpyxl
from shapely import wkb
from ftth_std import STD, LAY, LBL

ceil = lambda v: int(math.ceil(v - 1e-9))
SP = lambda: STD['spare']['splitters_closures']   # spare factor on splitters / closures (standards)

# ---------------------------------------------------------------- GPKG reader
def read_gpkg(path):
    con = sqlite3.connect(path); out = {}
    for (tn,) in con.execute("select table_name from gpkg_contents where data_type='features'"):
        gc = con.execute("select column_name from gpkg_geometry_columns where table_name=?", (tn,)).fetchone()[0]
        cols = [r[1] for r in con.execute('pragma table_info("%s")' % tn)]
        rows = []
        for r in con.execute('select * from "%s"' % tn):
            d = dict(zip(cols, r)); b = d.pop(gc)
            if b is None: d['geometry'] = None
            else:
                env = {0: 0, 1: 32, 2: 48, 3: 48, 4: 64}[(b[3] >> 1) & 7]
                d['geometry'] = wkb.loads(bytes(b[8 + env:]))
            rows.append(d)
        out[tn] = rows
    con.close(); return out

def L(layers, name): return layers.get(name, [])
def num(v, default=0.0):
    try: return float(v) if v not in (None, '') else default
    except Exception: return default

# ---------------------------------------------------------------- metrics f(gpkg)
def poles(layers):
    """pole estimate along main+branch cables: a pole every poles.span_m, plus corners > corner_deg, dedup within dedup_m (standards)"""
    pts, roles = [], []
    for nm in (LAY['main_cables'], LAY['branch_cables']):
        for r in L(layers, nm):
            g = r['geometry']
            if g is None: continue
            for ln in (g.geoms if g.geom_type.startswith('Multi') else [g]):
                c = np.array(ln.coords); Ln = ln.length; n = ceil(Ln / STD['poles']['span_m'])
                for k in range(n + 1):
                    p = ln.interpolate(min(Ln, k * Ln / n if n else 0)); pts.append((p.x, p.y)); roles.append('span')
                for i in range(1, len(c) - 1):
                    a = c[i, :2] - c[i - 1, :2]; b = c[i + 1, :2] - c[i, :2]; na, nb = np.hypot(*a), np.hypot(*b)
                    if na > 0 and nb > 0 and np.degrees(np.arccos(np.clip(a @ b / (na * nb), -1, 1))) > STD['poles']['corner_deg']: pts.append(tuple(c[i, :2])); roles.append('corner')
    special = []
    for nm in (LAY['center'], LAY['primary_splitters'], LAY['secondary_splitters'], LAY['closure_main'], LAY['closure_branch'], LAY['closure_merge']):
        for r in L(layers, nm):
            if r['geometry'] is not None: special.append((r['geometry'].x, r['geometry'].y))
    P = np.array(special + pts) if (special or pts) else np.zeros((0, 2)); R = ['special'] * len(special) + roles
    order = sorted(range(len(P)), key=lambda i: (0 if R[i] == 'special' else 1 if R[i] == 'corner' else 2))
    kp, keep = [], []
    for i in order:
        if kp and np.min(np.hypot(*(np.array(kp) - P[i]).T)) < STD['poles']['dedup_m']: continue
        kp.append(P[i]); keep.append(i)
    stay = sum(1 for i in keep if R[i] != 'span')
    return len(keep), stay, len(keep) - stay

def metrics(layers):
    m = {}
    M, B = L(layers, LAY['main_cables']), L(layers, LAY['branch_cables'])
    for tag, rows in (('M', M), ('B', B)):
        for r in rows: m['%s%d' % (tag, int(r['Cores']))] = m.get('%s%d' % (tag, int(r['Cores'])), 0.0) + num(r.get('Order_m'))
    m['Mlen'] = sum(num(r.get('Length_m')) for r in M); m['Blen'] = sum(num(r.get('Length_m')) for r in B)
    m['Mord'] = sum(num(r.get('Order_m')) for r in M); m['Bord'] = sum(num(r.get('Order_m')) for r in B)
    m['cab_all'] = sum(num(r.get('Order_m')) for r in M + B)
    P, S = L(layers, LAY['primary_splitters']), L(layers, LAY['secondary_splitters'])
    m['KP'], m['K'] = len(P), len(S)
    cl = {k: len(L(layers, n)) for k, n in (('cm', LAY['closure_main']), ('cb', LAY['closure_branch']), ('cd', LAY['closure_merge']))}
    m.update(cl); m['NC'] = sum(cl.values())
    m['PON'] = m['KP'] + STD['pon']['extra_ports']
    co = L(layers, LAY['center']); root = 0
    if co and co[0]['geometry'] is not None:
        for r in M:
            if r['geometry'] is not None and r['geometry'].distance(co[0]['geometry']) < STD['odf']['root_radius_m']: root += int(r['Cores'])
    m['ODF'] = ceil(root / STD['odf']['fibres_per_panel']) if root else 0
    # splices
    sp = 0
    for n in (LAY['closure_main'], LAY['closure_branch'], LAY['closure_merge']):
        sp += sum(num(r.get('Splices')) for r in L(layers, n))
    if any('Used_Ports' in r for r in P[:1]):
        sp += m['KP'] + sum(num(r.get('Used_Ports')) for r in P) + m['K'] + sum(num(r.get('Used_Ports')) for r in S) + m['PON']
    else:
        sp = 0   # Orem has no splice attributes: proportional rule below
    m['csp'] = sum(num(r.get('Splices')) for n in (LAY['closure_main'],LAY['closure_branch'],LAY['closure_merge']) for r in L(layers, n)) or 2.0 * m['NC']
    m['spl'] = sp if sp else (2.0 * m['NC'] + 2.0 * m['K'] + 2.0 * m['KP'])
    import ftth_poles
    sv = ftth_poles.survey(layers)
    if sv:
        c = ftth_poles.counts(sv); m.update(pol=c['total'], stay=c['stay'], mid=c['mid'], pol_new=c['new'], pol_existing=c['existing'], pol_branch=c['branch'], survey=1)
    else:
        m['pol'], m['stay'], m['mid'] = poles(layers)
    if LAY['buildings'] in layers: m['subs'] = len(layers[LAY['buildings']])
    elif P and 'Subscribers' in P[0]: m['subs'] = int(sum(num(r['Subscribers']) for r in P))
    else: m['subs'] = int(sum(num(r.get('Buildings')) for r in S))
    return m

# ---------------------------------------------------------------- workbook helpers
def label(ws, r):
    b = ws.cell(r, 2).value
    return str(b).split('\n')[0].split(' | ')[0].strip() if b else ''
def find_total(ws):
    for r in range(1, ws.max_row + 1):
        if str(ws.cell(r, 6).value or '').startswith('الإجمالي / Total') or str(ws.cell(r, 1).value or '').startswith('الإجمالي (قبل'): return r
def item_rows(ws):
    t = find_total(ws)
    return {r: label(ws, r) for r in range(5, t) if ws.cell(r, 4).value and ws.cell(r, 5).value is not None}

def rows_by_label(ws):
    """new-template layout: {row: first-line label} for item rows (qty col E, price col F)"""
    out = {}
    for r in range(5, ws.max_row + 1):
        b = ws.cell(r, 2).value
        if b and ws.cell(r, 4).value and ws.cell(r, 5).value is not None: out[r] = str(b).split('\n')[0].strip()
    return out
def find_row(rows, *keys):
    return next((r for r, t in rows.items() if all(k in t for k in keys)), None)

def table1_new(ws, mb, me):
    """current table: delta rule  new = old + (metric_new - metric_baseline); spare 20% applied on splitter/closure nets"""
    rows = rows_by_label(ws); q = {}
    d = lambda k: me.get(k, 0.0) - mb.get(k, 0.0)
    for r, t in rows.items():
        old = ws.cell(r, 5).value
        if not isinstance(old, (int, float)): continue
        m = re.match(r'(%s|%s) (\d+)' % (re.escape(LBL['cable_main']), re.escape(LBL['cable_branch'])), t); new = None
        if m: new = old + d(('M' if m.group(1) == LBL['cable_main'] else 'B') + m.group(2))
        elif t.startswith(LBL['primary_splitter']): new = ceil(me['KP'] * SP() - 1e-9)
        elif t.startswith(LBL['secondary_splitter']): new = ceil(me['K'] * SP() - 1e-9)
        elif t.startswith(LBL['closure']): new = ceil(me['NC'] * SP() - 1e-9)
        elif t.startswith(LBL['gpon_card']): new = ceil(me['PON'] / STD['pon']['ports_per_card'])
        elif t.startswith(LBL['odf']): new = max(old, ceil(me['ODF'] * 1.0)) if d('ODF') <= 0 else old + d('ODF')
        elif t.startswith(LBL['patch_cord']) or t.startswith(LBL['storage_cross']): new = ceil(me['KP'] * SP() - 1e-9)
        if new is not None: q[r] = max(0, int(round(new)))
    return q, []

def table2_new(ws1, ws2, q1, mb, me, subs0):
    """year-10 table: S2_new = S2_old + (S1_new - S1_old) for cables / splitters / closures; then patch cord + storage cross
    follow the year-10 primary splitters (+20%), GPON boards follow year-10 PON ports"""
    r1 = rows_by_label(ws1); r2 = rows_by_label(ws2); q = {}
    for r, t in r2.items():
        k = next((x for x, y in r1.items() if y == t), None)
        if k is None or k not in q1: continue
        old2 = ws2.cell(r, 5).value; old1 = ws1.cell(k, 5).value
        if isinstance(old2, (int, float)) and isinstance(old1, (int, float)) and (t.startswith(LBL['cable_main'].split()[0]) or t.startswith('سبلتر') or t.startswith('كلوجر')):
            q[r] = max(0, int(round(old2 + q1[k] - old1)))
    rP = find_row(r2, LBL['primary_splitter'])
    P2g = q.get(rP, ws2.cell(rP, 5).value) if rP else 0
    for r, t in r2.items():
        if t.startswith(LBL['patch_cord']) or t.startswith(LBL['storage_cross']): q[r] = P2g
        if t.startswith(LBL['gpon_card']):
            net = int(P2g // SP()) + STD['pon']['extra_ports']; q[r] = max(ws2.cell(r, 5).value or 0, ceil(net / STD['pon']['ports_per_card']))
        if t.startswith(LBL['odf']): q[r] = max(ws2.cell(r, 5).value or 0, q1.get(find_row(r1, LBL['odf']), 0))
    info = dict(SUBS0=subs0, SUBS10=ws2.parent['الملخص_المالي']['C4'].value if 'الملخص_المالي' in ws2.parent.sheetnames else subs0,
                P1=q1.get(find_row(r1, LBL['primary_splitter'])), P2=P2g)
    return q, info

def copy_style(src, dst):
    import copy as _c
    dst.font = _c.copy(src.font); dst.fill = _c.copy(src.fill); dst.border = _c.copy(src.border); dst.alignment = _c.copy(src.alignment); dst.number_format = src.number_format

def refresh_side_sheets(wb, layers, me, mb, q1, q2, info):
    # summary sheet
    if 'ملخص' in wb.sheetnames:
        s = wb['ملخص']; D = lambda key: me.get(key, 0) - mb.get(key, 0)
        cur = {s.cell(r, 1).value: s.cell(r, 2).value for r in range(2, s.max_row + 1)}
        vals = {'عدد المباني (مشتركين)': me['subs'], 'سبلتر رئيسي 1:8 (مع احتياط 20%)': q1.get(find_row(rows_by_label(wb['الكميات_الحالية']), LBL['primary_splitter'])), 'سبلتر فرعي 1:8 (مع احتياط 20%)': q1.get(find_row(rows_by_label(wb['الكميات_الحالية']), LBL['secondary_splitter'])), 'كلوجرات': q1.get(find_row(rows_by_label(wb['الكميات_الحالية']), 'كلوجر'))}
        for key, mk in (('طول الكوابل الرئيسية (م)', 'Mord'), ('طول الكوابل الفرعية (م)', 'Bord'), ('أعمدة (تقديري)', 'pol')):
            if isinstance(cur.get(key), (int, float)): vals[key] = int(round(cur[key] + D(mk)))
        if me.get('survey'): vals['أعمدة (تقديري)'] = me['pol']
        for r in range(2, s.max_row + 1):
            k = s.cell(r, 1).value
            if k in vals and not str(s.cell(r, 2).value).startswith('='): s.cell(r, 2).value = vals[k]
            if k == 'كلوجرات': s.cell(r, 3).value = 'الصافي %d + احتياط 20%%: %d تفرع رئيسي، %d تفرع فرعي، %d دمج/فصل' % (me['cm'] + me['cb'] + me['cd'], me['cm'], me['cb'], me['cd'])
            if k == 'أعمدة (تقديري)' and me.get('survey'):
                s.cell(r, 1).value = 'أعمدة (من المسح الميداني)'
                s.cell(r, 3).value = '%d مع شدّادة، %d وسطية. موجود %d (لا توريد)، جديد %d (يحتاج توريد)' % (me['stay'], me['mid'], me['pol_existing'], me['pol_new']) + ('' if me['pol'] == me['pol_existing'] + me['pol_new'] else '، غير محدد %d' % (me['pol'] - me['pol_existing'] - me['pol_new']))
            elif k == 'أعمدة (تقديري)':
                mm = re.findall(r'(\d+) مع شدّادة، (\d+) وسطية', str(s.cell(r, 3).value))
                if mm: s.cell(r, 3).value = '%d مع شدّادة، %d وسطية' % (int(mm[0][0]) + int(round(D('stay'))), int(mm[0][1]) + int(round(D('mid'))))
    # cable detail
    if 'تفصيل الكوابل' in wb.sheetnames:
        w = wb['تفصيل الكوابل']; tmpl = [w.cell(2, c) for c in range(1, 9)]
        rows = []
        for tag, nm, key in (('رئيسي', LAY['main_cables'], 'Run'), ('فرعي', LAY['branch_cables'], 'Primary_ID')):
            for r in L(layers, nm):
                rows.append([tag, r.get('Cable_ID'), r.get('Cores'), r.get('Used'), round((num(r.get('Cores')) - num(r.get('Used'))) / num(r.get('Cores'), 1) * 100, 1) if num(r.get('Cores')) else r.get('Spare_pct'), r.get('Length_m'), r.get('Order_m'), r.get(key) or r.get('Group_ID') or '-'])
        old = w.max_row
        for i, vals in enumerate(rows, 2):
            for c, v in enumerate(vals, 1):
                x = w.cell(i, c); x.value = v
                if i > old: copy_style(tmpl[c - 1], x)
        for i in range(len(rows) + 2, old + 1):
            for c in range(1, 9): w.cell(i, c).value = None
    # groups (Bleon)
    if 'المجموعات' in wb.sheetnames:
        g = wb['المجموعات']; P = sorted(L(layers, LAY['primary_splitters']), key=lambda r: r['Primary_ID']); S = L(layers, LAY['secondary_splitters']); BR = L(layers, LAY['branch_cables'])
        st = [g.cell(2, c) for c in range(1, 9)]; stt = [g.cell(g.max_row, c) for c in range(1, 9)]
        old = g.max_row
        for rr in range(2, old + 1):
            for c in range(1, 9): g.cell(rr, c).value = None
        own = [r for r in BR if r.get('Merged') != 'نعم']; mg = [r for r in BR if r.get('Merged') == 'نعم']
        k = 2
        for p in P:
            PK = lambda r: r.get('Primary_ID') or r.get('Group_ID'); o = [r for r in own if PK(r) == p['Primary_ID']]; mm = [r for r in mg if p['Primary_ID'] in str(r.get('Primary_ID')).split('+')]
            vals = [p['Primary_ID'], sum(1 for r in S if PK(r) == p['Primary_ID']), (int(num(p.get('Subscribers'))) if 'Subscribers' in p else int(sum(num(r.get('Buildings')) for r in S if PK(r) == p['Primary_ID']))), (int(num(p.get('Used_Ports'))) if 'Used_Ports' in p else sum(1 for r in S if PK(r) == p['Primary_ID'])), int(num(p.get('Spare_Ports'))), round(sum(num(r.get('Length_m')) for r in o), 1), '=ROUND(F%d*%s,1)' % (k, STD['cable_order_factor']), len(mm)]
            for c, v in enumerate(vals, 1): g.cell(k, c).value = v; copy_style(st[c - 1], g.cell(k, c))
            k += 1
        vals = ['المجموع'] + ['=SUM(%s2:%s%d)' % (c_, c_, k - 1) for c_ in 'BCDEFG'] + ['']
        for c, v in enumerate(vals, 1): g.cell(k, c).value = v; copy_style(st[c - 1], g.cell(k, c))
        k += 1
        vals = ['كبل مدمج (مشترك بين مجموعات)', '', '', '', '', round(sum(num(r.get('Length_m')) for r in mg), 1), '=ROUND(F%d*%s,1)' % (k, STD['cable_order_factor']), len(mg)]
        for c, v in enumerate(vals, 1): g.cell(k, c).value = v; copy_style(st[c - 1], g.cell(k, c))
        k += 1
        vals = ['الإجمالي الفرعي', '', '', '', '', '=F%d+F%d' % (k - 2, k - 1), '=G%d+G%d' % (k - 2, k - 1), '']
        for c, v in enumerate(vals, 1): g.cell(k, c).value = v; copy_style(st[c - 1], g.cell(k, c))

def apply_book(path_in, path_out, layers, mb, me, subs0):
    wb = openpyxl.load_workbook(path_in)
    s1, s2 = wb['الكميات_الحالية'], wb['الكميات_بعد_10_سنوات']
    q1, kept = table1_new(s1, mb, me)
    q2, info = table2_new(s1, s2, q1, mb, me, subs0)
    r1 = rows_by_label(s1)
    # hanging materials: STD['hanging_materials'].mode -> 'by_cable_length' (ceil(len/span/round)*round) or 'manual' (value stays: entered by the person doing the study)
    rh = find_row(r1, LBL['hanging'])
    if rh and STD['hanging_materials']['mode'] == 'by_cable_length':
        hb = STD['hanging_by_cable_length']; hang = lambda tot: ceil(tot / hb['span_m'] / hb['round_to']) * hb['round_to']
        cab = lambda ws, qq: sum((qq.get(r, ws.cell(r, 5).value) or 0) for r, t in rows_by_label(ws).items() if t.startswith('كابل'))
        q1[rh] = hang(cab(s1, q1))
        rh2 = find_row(rows_by_label(s2), LBL['hanging']); q2[rh2] = hang(cab(s2, q2))
    for r, v in q1.items(): s1.cell(r, 5).value = v
    for r, v in q2.items(): s2.cell(r, 5).value = v
    # financial summary and side sheet link by formula to the spare-inclusive quantity cells
    refresh_side_sheets(wb, layers, me, mb, q1, q2, {'SUBS0': info['SUBS0'], 'SUBS10': info['SUBS10'], 'NEWSEC': 0, 'NEWPRIM': 0, 'nclo_m': 0, 'nclo_b': 0, 'cab_m': 0})
    wb.save(path_out)
    return q1, q2, info, kept

# ---------------------------------------------------------------- main
def _opt(name):
    try: return __import__(name)
    except Exception as e:
        print('تنبيه: تعذّر تحميل %s (%s) - تُتخطى المخرجات المرتبطة به.' % (name, e)); return None

def change_report(tpl_path, q1, mb, me, out_path):
    """quantities before/after per item + $ delta (company BOQ, current table)"""
    wb = openpyxl.load_workbook(tpl_path); ws = wb['الكميات_الحالية']; rows = rows_by_label(ws); lines = []
    tot = 0.0
    for r, t in rows.items():
        old = ws.cell(r, 5).value; price = ws.cell(r, 6).value
        if r not in q1 or not isinstance(old, (int, float)) or not isinstance(price, (int, float)): continue
        new = q1[r]
        if abs(new - old) < 1e-9: continue
        d = (new - old) * price; tot += d; lines.append((t, old, new, new - old, price, round(d, 2)))
    newitems = []
    have = set(rows.values())
    for k, v in me.items():
        m = re.match(r'^([MB])(\d+)$', k)
        if m and v > 0 and mb.get(k, 0) <= 0 and ('كابل %s %s ألياف' % ('رئيسي' if m.group(1) == 'M' else 'فرعي', m.group(2))) not in have: newitems.append('كبل %s %s ألياف: %d م (بند جديد، أضفه يدوياً مع سعره)' % ('رئيسي' if m.group(1) == 'M' else 'فرعي', m.group(2), round(v)))
    from openpyxl.styles import Font, PatternFill, Alignment
    w = openpyxl.Workbook(); s = w.active; s.title = 'تقرير التغييرات'; s.sheet_view.rightToLeft = True
    s.append(['البند', 'الكمية السابقة', 'الكمية الجديدة', 'الفرق', 'سعر الوحدة (USD)', 'الفرق بالدولار'])
    for c in s[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='14284B'); c.alignment = Alignment(horizontal='center')
    for l in sorted(lines, key=lambda x: -abs(x[5])): s.append(list(l))
    s.append(['الفرق الإجمالي (قبل احتياطي الطوارئ %d%%)' % round(STD['contingency'] * 100) + '', '', '', '', '', round(tot, 2)]); s.append(['الفرق الإجمالي بعد الاحتياطي', '', '', '', '', round(tot * (1 + STD['contingency']), 2)])
    for c in s[s.max_row]: c.font = Font(bold=True)
    for c in s[s.max_row - 1]: c.font = Font(bold=True)
    for n in newitems: s.append([n])
    for col, wd in zip('ABCDEF', (46, 14, 14, 10, 16, 16)): s.column_dimensions[col].width = wd
    w.save(out_path)
    return lines, tot, newitems

def main():
    import ftth_std
    argv = sys.argv[1:]; study_dir = None
    if '--study' in argv:
        i = argv.index('--study'); study_dir = argv[i + 1]; del argv[i:i + 2]
    poles_arg = None
    if '--poles' in argv:
        i = argv.index('--poles'); poles_arg = argv[i + 1]; del argv[i:i + 2]
    here = os.path.abspath(study_dir or os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..'))
    cfg = json.load(open(os.path.join(here, 'study.json'), encoding='utf-8'))
    package_dir = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), '..'))
    ftth_std.load(cfg, package_dir); OUT = STD['outputs']
    import ftth_ar; _con = ftth_ar.install(STD.get('console', {}).get('arabic_shaping', 'auto'))
    args = [a for a in argv if not a.startswith('--')]; force = '--force' in argv
    gp = args[0] if args else os.path.join(here, '..', cfg['gpkg'])
    gp = os.path.abspath(gp)
    print('قراءة الملف:', gp)
    layers = read_gpkg(gp)
    import ftth_poles
    psrc = ftth_poles.attach(layers, gp, here, cfg, poles_arg)
    print('الأعمدة: ' + ('من المسح الميداني (%s): %d عمود' % (psrc, len(layers[LAY['poles']])) if psrc else 'تقدير (لا توجد طبقة أعمدة من المسح)'))
    me = metrics(layers); mb = cfg['baseline']
    stamp = datetime.datetime.now().strftime('%Y%m%d_%H%M')
    out = os.path.join(here, 'نتائج_التحديث', stamp); os.makedirs(out, exist_ok=True)
    _tee = open(os.path.join(out, 'run_summary.txt'), 'w', encoding='utf-8'); _con.set_tee(_tee); _so = _con.s
    ex = _opt('ftth_extra')
    if ex and not OUT.get('loss_budget_xlsx', True): ex.loss_xlsx = lambda *a, **k: None
    # 1) validation
    if ex:
        issues = ex.validate(layers)
        errs = [m for l, m in issues if l == 'خطأ']
        print('\n== فحص الملف ==')
        for l, m in issues: print(' [%s] %s' % (l, m))
        if not issues: print(' لا ملاحظات.')
        if errs and not force:
            print('\nتوقف التحديث بسبب الأخطاء أعلاه. أصلحها في QGIS ثم أعد التشغيل (أو شغّل مع --force للمتابعة رغم الأخطاء).')
            _con.set_tee(None); _tee.close(); print('Summary file (UTF-8, readable): ' + os.path.join(out, 'run_summary.txt'))
            return
    # 2) loss budget
    ltxt = None; lsum = None
    if ex:
        rows, bad = ex.loss_budget(layers); lsum = ex.loss_summary(rows)
        if lsum:
            ex.loss_xlsx(rows, bad, os.path.join(out, 'ميزانية_الفقد_%s.xlsx' % cfg.get('file_town', cfg['study'])), cfg.get('town_ar', cfg['study']))
            ltxt = 'المحسوب: متوسط %.2f dB (أقصى %.2f، مع الهامش %.2f) - الهدف: %s' % (lsum['avg'], lsum['max'], lsum['maxm'], STD['loss']['target_text'])
            print('\n== ميزانية الفقد == متوسط %.2f dB، أقصى %.2f dB، أقصى مع هامش 3 dB: %.2f (الحد %d)' % (lsum['avg'], lsum['max'], lsum['maxm'], ex.LOSS['limit']))
            if lsum['avg'] >= STD['loss']['target_avg_db']: print(' تحذير: متوسط الفقد %.2f فوق الهدف %g dB' % (lsum['avg'], STD['loss']['target_avg_db']))
            if lsum['over']: print(' تحذير: %d سبلتر فرعي يتجاوز %d dB مع الهامش (انظر ملف ميزانية الفقد)' % (lsum['over'], ex.LOSS['limit']))
            elif lsum['over28']: print(' ملاحظة: %d سبلتر فرعي يتجاوز 28 dB (Class B+) مع الهامش، وجميعها ضمن 32 dB (Class C+).' % lsum['over28'])
            if bad: print(' تحذير: %d سبلتر فرعي غير موصول بشبكة الكبلات: %s' % (len(bad), '، '.join(map(str, bad[:8]))))
    # 3) BOQ
    if STD['hanging_materials']['mode'] != 'by_cable_length': print('مواد التعليق: تُدخل يدوياً في قالب BOQ ولا تغيّرها الأداة (Hanging materials = manual input).')
    generic = bool(cfg.get('generic'))
    prices_all = {}
    pf = os.path.join(here, 'الأسعار.json')
    if os.path.exists(pf):
        try: prices_all = json.load(open(pf, encoding='utf-8'))
        except Exception as e: print('تنبيه: تعذّرت قراءة الأسعار.json:', e)
    s1 = openpyxl.load_workbook(os.path.join(here, cfg['templates'][0]))['الكميات_الحالية']
    have = set(rows_by_label(s1).values())
    for k, v in me.items():
        m = re.match(r'^([MB])(\d+)$', k)
        if m and v > 0 and mb.get(k, 0) <= 0 and ('كابل %s %s ألياف' % ('رئيسي' if m.group(1) == 'M' else 'فرعي', m.group(2))) not in have:
            print('تنبيه: ظهر كبل %s بعدد ألياف %s ولا يوجد له بند في الجدول، أضفه يدوياً.' % ('رئيسي' if m.group(1) == 'M' else 'فرعي', m.group(2)))
    comp_path = comp_tpl = comp_q1 = None; gen_priced = 0
    for tpl in cfg['templates']:
        dst = os.path.join(out, tpl.replace('.xlsx', '_%s.xlsx' % stamp))
        if generic:
            import ftth_boq
            q1, q2, gm = ftth_boq.refresh_workbook(os.path.join(here, tpl), dst, me, STD, cfg.get('growth'))
            kind = 'company' if 'الشركة' in tpl else 'real'
            try:
                if openpyxl.load_workbook(dst)['مدخلات']['B3'].value in (None, ''): print('  تنبيه: أشهر العمل غير مُدخلة (صفحة مدخلات في القالب). تُدخل بعد خطة التنفيذ.')
            except Exception: pass
            pr = {k: v for k, v in (prices_all.get(kind) or {}).items() if isinstance(v, (int, float))}
            if pr: print('  أسعار %s: طُبّقت على %d بند' % (kind, ftth_boq.apply_prices(dst, dst, pr)))
            if 'الشركة' in tpl: gen_priced = ftth_boq.priced_total(dst)
        else:
            q1, q2, info, kept = apply_book(os.path.join(here, tpl), dst, layers, mb, me, me['subs'])
        print('تم:', os.path.basename(dst))
        if 'الشركة' in tpl: comp_path, comp_tpl, comp_q1 = dst, os.path.join(here, tpl), (None if generic else q1)
    print('مع احتياط 20%%: سبلترات رئيسية %d، فرعية %d، كلوجرات %d' % tuple(ceil(x * SP() - 1e-9) for x in (me['KP'], me['K'], me['NC'])))
    print('أطوال الكبلات المطلوبة (م): رئيسية %d، فرعية %d' % (sum(v for k, v in me.items() if re.match(r'^M\d+$', k)), sum(v for k, v in me.items() if re.match(r'^B\d+$', k))))
    # 4) change report
    if comp_tpl and comp_q1 is not None and ex and OUT.get('change_report', True):
        lines, tot, newitems = change_report(comp_tpl, comp_q1, mb, me, os.path.join(out, 'تقرير_التغييرات_%s.xlsx' % stamp))
        print('\n== تقرير التغييرات (مقارنة بالمسلَّم) ==')
        for t, o, n, d, p, dd in sorted(lines, key=lambda x: -abs(x[5]))[:12]: print(' %-34s %9s -> %-9s (%+g) %+10.2f $' % (t[:34], format(o, 'g'), format(n, 'g'), d, dd))
        if not lines: print(' لا تغيير في الكميات.')
        print(' الفرق الإجمالي: %+.2f $ (قبل الاحتياطي) ، %+.2f $ (بعد 5%%)' % (tot, tot * (1 + STD['contingency'])))
        for n in newitems: print(' بند جديد:', n)
    # 5) company workbook details + client files
    if comp_path:
        try:
            wb = openpyxl.load_workbook(comp_path)
            if 'بيانات العرض' in wb.sheetnames:
                bd = wb['بيانات العرض']; bd['C11'] = datetime.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
                if ltxt: bd['C22'] = ltxt
            wb.save(comp_path)
        except Exception as e: print('تنبيه: تعذّر تحديث بيانات العرض:', e)
        of = _opt('ftth_offer')
        if generic and not gen_priced:
            print('عرض السعر: لم تُدخل الأسعار بعد (الأسعار.json أو الخلايا الصفراء في القالب). أُنتجت الكميات فقط.'); of = None
        if of:
            town = cfg.get('town_ar', cfg['study']); fl = cfg.get('file_town', cfg['study'])
            try:
                if OUT.get('client_xlsx', True): of.client_xlsx(comp_path, os.path.join(out, 'FTTH_%s_عرض_سعر.xlsx' % fl), town); print('تم: FTTH_%s_عرض_سعر.xlsx' % fl)
                d = of.offer_data(openpyxl.load_workbook(comp_path), town)
                if OUT.get('client_docx', True): of.build_docx(d, os.path.join(out, 'عرض_سعر_FTTH_%s.docx' % fl)); print('تم: عرض_سعر_FTTH_%s.docx (الإجمالي %s USD)' % (fl, format(d['fin'], ',')))
            except ImportError as e: print('تنبيه: مكتبة ناقصة لملف الورد (%s). شغّل: pip install python-docx' % e)
            except Exception as e: print('تنبيه: تعذّر توليد عرض السعر:', e)
    # 6) link table + KML
    if ex and OUT.get('link_table', True):
        try: ex.link_table(layers, os.path.join(out, 'جدول_ارتباط_السبلترات_%s.xlsx' % cfg.get('file_town', cfg['study']))); print('تم: جدول ارتباط السبلترات')
        except Exception as e: print('تنبيه: تعذّر جدول الارتباط:', e)
    if layers.get(LAY['poles']) and OUT.get('poles_audit_xlsx', True):
        try: ftth_poles.audit_xlsx(layers, os.path.join(out, 'تدقيق_الأعمدة_%s.xlsx' % cfg.get('file_town', cfg['study'])), cfg.get('town_ar', cfg['study']), psrc or ''); print('تم: تدقيق الأعمدة (أعمدة جديدة %d تحتاج توريداً، موجودة %d)' % (me['pol_new'], me['pol_existing']))
        except Exception as e: print('تنبيه: تعذّر تدقيق الأعمدة:', e)
    km = _opt('ftth_kml') if OUT.get('kml', True) else None
    if km:
        try: n = km.build_all(layers, cfg.get('town_ar', cfg['study']), out, cfg.get('file_town', cfg['study'])); print('تم: ملفات Google Earth (%d قطاع)' % n)
        except Exception as e: print('تنبيه: تعذّر توليد KML:', e)
    print('بنود ثابتة لا تتغير بالتحديث: غرفة المقسم، الطاقة، المتفرقات، الأشهر، المحروقات، الحفر وغرف التفتيش (تُدخل يدوياً عند وجودها).')
    print('الملفات في:', out)
    _con.set_tee(None); _tee.close(); print('Summary file (UTF-8, readable): ' + os.path.join(out, 'run_summary.txt'))

if __name__ == '__main__': main()
