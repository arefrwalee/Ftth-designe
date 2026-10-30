# -*- coding: utf-8 -*-
"""Phase 2: surveyed poles layer.
Sources (first found wins): --poles <file> on the command line, `poles_file` in study.json, a poles layer inside the study gpkg.
File formats: GeoPackage (point layer) or CSV (lon/lat WGS84, or x/y in the study UTM zone).
When a survey exists the tool counts poles from it (existing poles are not supplied, new poles are); otherwise the old estimate
(every poles.span_m + corners, dedup) stays in use. Checks: duplicate ids, unknown type/role, cable without a nearby pole, span longer than max_span_m."""
import os, re, csv, math, sqlite3, struct
import numpy as np
from shapely import wkb
from shapely.geometry import Point
import shapely
from ftth_std import STD, LAY

# ---------------------------------------------------------------- coordinates
def lonlat_to_utm(lon, lat, zone=None):
    """WGS84 lon/lat (deg) -> UTM easting/northing (northern hemisphere), inverse of ftth_kml.to_lonlat"""
    zone = zone or STD['crs']['utm_zone']
    a = 6378137.0; f = 1 / 298.257223563; k0 = 0.9996; e2 = f * (2 - f); ep2 = e2 / (1 - e2)
    phi = math.radians(lat); lam = math.radians(lon); lam0 = math.radians(zone * 6 - 183)
    N = a / math.sqrt(1 - e2 * math.sin(phi) ** 2); T = math.tan(phi) ** 2; C = ep2 * math.cos(phi) ** 2; A = (lam - lam0) * math.cos(phi)
    M = a * ((1 - e2 / 4 - 3 * e2 ** 2 / 64 - 5 * e2 ** 3 / 256) * phi - (3 * e2 / 8 + 3 * e2 ** 2 / 32 + 45 * e2 ** 3 / 1024) * math.sin(2 * phi)
             + (15 * e2 ** 2 / 256 + 45 * e2 ** 3 / 1024) * math.sin(4 * phi) - (35 * e2 ** 3 / 3072) * math.sin(6 * phi))
    x = k0 * N * (A + (1 - T + C) * A ** 3 / 6 + (5 - 18 * T + T ** 2 + 72 * C - 58 * ep2) * A ** 5 / 120) + 500000.0
    y = k0 * (M + N * math.tan(phi) * (A ** 2 / 2 + (5 - T + 9 * C + 4 * C ** 2) * A ** 4 / 24 + (61 - 58 * T + T ** 2 + 600 * C - 330 * ep2) * A ** 6 / 720))
    return x, y

# ---------------------------------------------------------------- text normalisation
def norm(s):
    s = str(s if s is not None else '').strip().lower()
    s = re.sub('[ً-ٰٟـ]', '', s)
    return s.translate(str.maketrans({'أ': 'ا', 'إ': 'ا', 'آ': 'ا', 'ة': 'ه', 'ى': 'ي', '_': '', ' ': '', '-': ''}))

def _pick(row, aliases):
    low = {norm(k): k for k in row}
    for a in aliases:
        k = low.get(norm(a))
        if k is not None and row[k] not in (None, ''): return row[k]
    return None

def _in(value, lists):
    v = norm(value)
    for key, vals in lists.items():
        if v in [norm(x) for x in vals]: return key
    return None

# ---------------------------------------------------------------- readers
def _gpkg_srs(path, table):
    con = sqlite3.connect(path)
    try: r = con.execute('select srs_id from gpkg_geometry_columns where table_name=?', (table,)).fetchone()
    finally: con.close()
    return r[0] if r else None

def _to_study_crs(g, srs):
    if g is None or srs != 4326: return g
    x, y = lonlat_to_utm(g.x, g.y); return Point(x, y)

def read_csv(path):
    raw = open(path, 'rb').read()
    for enc in ('utf-8-sig', 'cp1256'):
        try: text = raw.decode(enc); break
        except UnicodeDecodeError: continue
    delim = ';' if text.split('\n', 1)[0].count(';') > text.split('\n', 1)[0].count(',') else ','
    rows = list(csv.DictReader(text.splitlines(), delimiter=delim)); F = STD['poles']['survey']['fields']; out = []
    for r in rows:
        lon, lat = _pick(r, F['lon']), _pick(r, F['lat']); x, y = _pick(r, F['x']), _pick(r, F['y']); g = None
        try:
            if lon is not None and lat is not None: g = Point(*lonlat_to_utm(float(str(lon).replace(',', '.')), float(str(lat).replace(',', '.'))))
            elif x is not None and y is not None: g = Point(float(str(x).replace(',', '.')), float(str(y).replace(',', '.')))
        except ValueError: g = None
        d = dict(r); d['geometry'] = g; out.append(d)
    return out

def read_poles_file(path):
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.txt'): return read_csv(path)
    if ext == '.gpkg':
        from ftth_update import read_gpkg
        L = read_gpkg(path); name = _find_layer_name(L)
        if name is None: raise ValueError('لا توجد طبقة أعمدة (نقاط) في %s' % os.path.basename(path))
        srs = _gpkg_srs(path, name)
        return [dict(r, geometry=_to_study_crs(r['geometry'], srs)) for r in L[name]]
    raise ValueError('صيغة ملف الأعمدة غير مدعومة: ' + ext)

def _find_layer_name(layers):
    S = STD['poles']['survey']
    if S['layer'] in layers: return S['layer']
    for n, rows in layers.items():                     # fallback: a point layer that carries a pole-id field
        if rows and rows[0].get('geometry') is not None and rows[0]['geometry'].geom_type == 'Point' and _pick(rows[0], S['fields']['id']) is not None and n not in [v for k, v in LAY.items() if k != 'poles']:
            if any(norm(k) in [norm(a) for a in S['fields']['id']] for k in rows[0]): return n
    return None

# ---------------------------------------------------------------- normalise + count
def normalise(rows):
    S = STD['poles']['survey']; out = []
    for i, r in enumerate(rows, 1):
        t = _in(_pick(r, S['fields']['type']), S['type_values']) or 'unknown'
        ro = _in(_pick(r, S['fields']['role']), S['role_values']) or 'unknown'
        h = _pick(r, S['fields']['height'])
        out.append(dict(id=str(_pick(r, S['fields']['id']) or 'P?%d' % i), type=t, role=ro, height=h, geometry=r.get('geometry')))
    return out

def counts(P):
    c = dict(total=len(P))
    for k in ('existing', 'new'): c[k] = sum(1 for p in P if p['type'] == k)
    c['unknown_type'] = c['total'] - c['existing'] - c['new']
    c['stay'] = sum(1 for p in P if p['role'] in ('stay', 'branch')); c['mid'] = sum(1 for p in P if p['role'] == 'mid')
    c['branch'] = sum(1 for p in P if p['role'] == 'branch'); c['unknown_role'] = c['total'] - c['stay'] - c['mid']
    return c

def attach(layers, gpkg_path, study_dir, cfg, cli_file=None):
    """put the surveyed poles into layers[LAY['poles']] (normalised rows). Returns a source description or None."""
    f = cli_file or cfg.get('poles_file')
    if f:
        p = f if os.path.isabs(f) else os.path.abspath(os.path.join(study_dir, '..', f))
        if not os.path.exists(p): p2 = os.path.abspath(os.path.join(study_dir, f)); p = p2 if os.path.exists(p2) else p
        if not os.path.exists(p): raise FileNotFoundError('ملف الأعمدة غير موجود: ' + p)
        rows = read_poles_file(p); src = os.path.basename(p)
    else:
        name = _find_layer_name(layers)
        if name is None: return None
        srs = _gpkg_srs(gpkg_path, name); rows = [dict(r, geometry=_to_study_crs(r['geometry'], srs)) for r in layers[name]]; src = 'طبقة %s في %s' % (name, os.path.basename(gpkg_path))
    layers[LAY['poles']] = normalise(rows); return src

def suggest(cable_lines, special_pts):
    """أعمدة مقترحة بالطريقة التقديرية نفسها (عمود كل span_m + الزوايا + إزالة المكرر ضمن dedup_m).
    cable_lines: [LineString]; special_pts: [(x,y)] مواقع المركز والسبلترات والكلوجرات (تُحسب أعمدة شدّادة أولاً).
    يرجع [(x, y, 'stay'|'mid')]."""
    S = STD['poles']; pts, roles = [], []
    for ln in cable_lines:
        c = np.array(ln.coords); Ln = ln.length; n = max(1, math.ceil(Ln / S['span_m']))
        for k in range(n + 1):
            q = ln.interpolate(min(Ln, k * Ln / n)); pts.append((q.x, q.y)); roles.append('span')
        for i in range(1, len(c) - 1):
            a = c[i, :2] - c[i - 1, :2]; b = c[i + 1, :2] - c[i, :2]; na, nb = np.hypot(*a), np.hypot(*b)
            if na > 0 and nb > 0 and np.degrees(np.arccos(np.clip(a @ b / (na * nb), -1, 1))) > S['corner_deg']:
                pts.append(tuple(c[i, :2])); roles.append('corner')
    P = np.array(list(special_pts) + pts) if (special_pts or pts) else np.zeros((0, 2))
    R = ['special'] * len(special_pts) + roles
    order = sorted(range(len(P)), key=lambda i: (0 if R[i] == 'special' else 1 if R[i] == 'corner' else 2))
    kp, out = [], []
    for i in order:
        if kp and np.min(np.hypot(*(np.array(kp) - P[i]).T)) < S['dedup_m']: continue
        kp.append(P[i]); out.append((float(P[i][0]), float(P[i][1]), 'mid' if R[i] == 'span' else 'stay'))
    return out

def survey(layers):
    P = layers.get(LAY['poles'])
    if P and 'role' not in P[0]: layers[LAY['poles']] = P = normalise(P)    # raw gpkg layer not yet normalised
    return P if P else None

# ---------------------------------------------------------------- audit
def _cables(layers):
    for tag, key in (('رئيسي', LAY['main_cables']), ('فرعي', LAY['branch_cables'])):
        for r in layers.get(key, []):
            g = r['geometry']
            if g is None: continue
            yield tag, str(r.get('Cable_ID')), g

def audit(layers):
    """returns (issues, per_cable_rows, per_pole_rows)"""
    P = layers.get(LAY['poles']) or []; S = STD['poles']['survey']; tol = S['pole_to_cable_m']; mx = S['max_span_m']
    pts = [p for p in P if p['geometry'] is not None]; arr = np.array([[p['geometry'].x, p['geometry'].y] for p in pts]) if pts else np.zeros((0, 2))
    pg = shapely.points(arr) if len(pts) else np.array([])
    cab_rows = []; no_pole = []; long_span = []
    near_cable = {i: (None, 1e18) for i in range(len(pts))}
    for tag, cid, g in _cables(layers):
        parts = list(g.geoms) if g.geom_type.startswith('Multi') else [g]
        for ln in parts:
            d = shapely.distance(ln, pg) if len(pts) else np.array([])
            for i in range(len(pts)):
                if d[i] < near_cable[i][1]: near_cable[i] = (cid, float(d[i]))
            idx = np.where(d <= tol)[0] if len(pts) else []
            ch = sorted(float(shapely.line_locate_point(ln, pg[i])) for i in idx)
            gaps = np.diff([0.0] + ch + [ln.length]) if ch else np.array([ln.length])
            mg = float(gaps.max()) if len(gaps) else 0.0
            flag = ''
            if not len(idx): flag = 'بلا عمود'; no_pole.append(cid)
            elif mg > mx: flag = 'باع أطول من الحد'; long_span.append((mg, cid))
            cab_rows.append([cid, tag, round(ln.length, 1), int(len(idx)), round(mg, 1), flag])
    issues = []
    ids = [p['id'] for p in P]; dup = sorted({i for i in ids if ids.count(i) > 1})
    if dup: issues.append(('تنبيه', 'معرّفات أعمدة مكررة: %s' % '، '.join(dup[:8]) + (' …' if len(dup) > 8 else '')))
    c = counts(P)
    if c['unknown_type']: issues.append(('تنبيه', '%d عمود بلا نوع معروف (موجود/جديد) فلا يُحسب في التوريد ولا في الموجود' % c['unknown_type']))
    if c['unknown_role']: issues.append(('تنبيه', '%d عمود بلا دور معروف (شدّادة/وسطي/تفرع)' % c['unknown_role']))
    nog = [p['id'] for p in P if p['geometry'] is None]
    if nog: issues.append(('تنبيه', '%d عمود بلا إحداثيات (يُتجاهل): %s' % (len(nog), '، '.join(nog[:6]))))
    if len(pts) > 1:
        from shapely.strtree import STRtree
        t = STRtree(pg); close = 0
        for i, g in enumerate(pg):
            for j in t.query(g.buffer(S['duplicate_m'])):
                if j > i and g.distance(pg[j]) <= S['duplicate_m']: close += 1
        if close: issues.append(('تنبيه', '%d زوج أعمدة أقرب من %g م (مكرر محتمل)' % (close, S['duplicate_m'])))
    if no_pole: issues.append(('تنبيه', '%d كبل بلا عمود قريب (ضمن %g م): %s' % (len(no_pole), tol, '، '.join(no_pole[:8]) + (' …' if len(no_pole) > 8 else ''))))
    if long_span:
        long_span.sort(reverse=True)
        issues.append(('تنبيه', '%d كبل فيه باع أطول من %g م (الأطول %.0f م في %s)' % (len(long_span), mx, long_span[0][0], long_span[0][1])))
    pole_rows = [[p['id'], {'existing': 'موجود', 'new': 'جديد'}.get(p['type'], 'غير محدد'), {'stay': 'شدّادة', 'mid': 'وسطي', 'branch': 'تفرع'}.get(p['role'], 'غير محدد'), p['height'],
                  near_cable[i][0] if near_cable[i][1] < 1e17 else '', round(near_cable[i][1], 1) if near_cable[i][1] < 1e17 else ''] for i, p in enumerate(pts)]
    return issues, cab_rows, pole_rows

def validate(layers):
    if not layers.get(LAY['poles']): return []
    return audit(layers)[0]

def audit_xlsx(layers, out_path, town, source):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    issues, cab, pol = audit(layers); c = counts(layers[LAY['poles']]); S = STD['poles']['survey']
    acc = STD['offer']['design']['accent']
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'ملخص الأعمدة'; ws.sheet_view.rightToLeft = True
    ws.append(['تدقيق طبقة الأعمدة - %s' % town]); ws['A1'].font = Font(bold=True, size=13)
    ws.append(['المصدر', source]); ws.append([])
    for a, b in (('إجمالي الأعمدة', c['total']), ('موجود (لا يُحسب له توريد)', c['existing']), ('جديد (يحتاج توريد)', c['new']), ('نوع غير محدد', c['unknown_type']),
                 ('شدّادة + تفرع', c['stay']), ('منها تفرع', c['branch']), ('وسطي', c['mid']), ('دور غير محدد', c['unknown_role']),
                 ('أقصى باع مسموح (م)', S['max_span_m']), ('مسافة العمود من الكبل لاعتباره قريباً (م)', S['pole_to_cable_m'])): ws.append([a, b])
    ws.append([]); ws.append(['الملاحظات'])
    for lv, m in issues: ws.append(['[%s]' % lv, m])
    if not issues: ws.append(['', 'لا ملاحظات'])
    ws.column_dimensions['A'].width = 44; ws.column_dimensions['B'].width = 70
    def sheet(title, head, rows, widths):
        w = wb.create_sheet(title); w.sheet_view.rightToLeft = True; w.append(head)
        for cl in w[1]: cl.font = Font(bold=True, color='FFFFFF'); cl.fill = PatternFill('solid', fgColor=acc); cl.alignment = Alignment(horizontal='center')
        for r in rows: w.append(r)
        for col, wd in zip('ABCDEF', widths): w.column_dimensions[col].width = wd
        w.freeze_panes = 'A2'
    sheet('الكبلات', ['الكبل', 'النوع', 'الطول (م)', 'أعمدة قريبة', 'أطول باع (م)', 'ملاحظة'], sorted(cab, key=lambda r: (r[5] == '', -r[4])), (16, 10, 12, 12, 14, 22))
    sheet('الأعمدة', ['العمود', 'النوع', 'الدور', 'الارتفاع', 'أقرب كبل', 'المسافة (م)'], pol, (16, 12, 12, 12, 16, 12))
    wb.save(out_path)
