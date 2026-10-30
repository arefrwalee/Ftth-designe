# -*- coding: utf-8 -*-
"""Client offer files from the company BOQ workbook: Excel offer + Word offer (python-docx, no PDF)."""
import datetime, copy
import openpyxl
from openpyxl.styles import PatternFill, Border, Font

from ftth_std import STD
ONES = ['', 'واحد', 'اثنان', 'ثلاثة', 'أربعة', 'خمسة', 'ستة', 'سبعة', 'ثمانية', 'تسعة', 'عشرة', 'أحد عشر', 'اثنا عشر', 'ثلاثة عشر', 'أربعة عشر', 'خمسة عشر', 'ستة عشر', 'سبعة عشر', 'ثمانية عشر', 'تسعة عشر']
TENS = ['', '', 'عشرون', 'ثلاثون', 'أربعون', 'خمسون', 'ستون', 'سبعون', 'ثمانون', 'تسعون']
HUND = ['', 'مائة', 'مائتان', 'ثلاثمائة', 'أربعمائة', 'خمسمائة', 'ستمائة', 'سبعمائة', 'ثمانمائة', 'تسعمائة']
def _b1000(n):
    p = []
    if n >= 100: p.append(HUND[n // 100]); n %= 100
    if n:
        if n < 20: p.append(ONES[n])
        else:
            u = n % 10; p.append((ONES[u] + ' و' + TENS[n // 10]) if u else TENS[n // 10])
    return ' و'.join(p)
def words(n):
    n = int(n); parts = []
    m, rest = divmod(n, 1000000); th, rest = divmod(rest, 1000)
    if m: parts.append('مليون' if m == 1 else 'مليونان' if m == 2 else _b1000(m) + (' ملايين' if m <= 10 else ' مليوناً'))
    if th: parts.append('ألف' if th == 1 else 'ألفان' if th == 2 else _b1000(th) + (' آلاف' if th <= 10 else ' ألفاً'))
    if rest: parts.append(_b1000(rest))
    return ' و'.join(parts) if parts else 'صفر'

# ------------------------------------------------------------------ client Excel
def client_xlsx(company_xlsx, out_path, town, loss_text=None):
    wb = openpyxl.load_workbook(company_xlsx)
    keep = ['بيانات العرض', 'الملخص_المالي', 'الكميات_الحالية']
    for n in list(wb.sheetnames):
        if n not in keep: del wb[n]
    wb._sheets = [wb[n] for n in keep]; wb.active = 0
    for ws in wb: ws.sheet_view.tabSelected = (ws.title == keep[0])
    sm = wb['الملخص_المالي']
    for m in list(sm.merged_cells.ranges): sm.unmerge_cells(str(m))
    for r in range(1, sm.max_row + 1):
        for c in (3, 4):
            cell = sm.cell(r, c); cell.value = None; cell.fill = PatternFill(fill_type=None); cell.border = Border()
    for r in (1, 2, 15): sm.merge_cells('A%d:B%d' % (r, r))
    sm['A2'] = 'الأرقام تُحدَّث تلقائياً من جدول الكميات.'
    sm['A15'] = 'كل البنود مسعّرة بالدولار الأمريكي (USD).'; sm['A15'].font = Font(name='Arial', sz=9, i=True, color='595959')
    sm.column_dimensions['C'].width = 4; sm.column_dimensions['D'].width = 4
    s1 = wb['الكميات_الحالية']; s1['A1'] = 'جدول الكميات والأسعار - شبكة FTTH الهوائية - بلدة %s' % town
    wb.save(out_path)

# ------------------------------------------------------------------ data from company BOQ
def offer_data(wb, town, loss_text=None):
    ws = wb['الكميات_الحالية']; bd = wb['بيانات العرض'] if 'بيانات العرض' in wb.sheetnames else None
    rows = []
    import re as _re
    def qv(r, depth=0):
        v = ws.cell(r, 5).value
        m = _re.match(r'^=E(\d+)$', str(v or ''))
        if m and depth < 5: return qv(int(m.group(1)), depth + 1)
        return v if isinstance(v, (int, float)) else 0
    for r in range(5, ws.max_row + 1):
        b = ws.cell(r, 2).value; q = qv(r); p = ws.cell(r, 6).value
        if b and ws.cell(r, 4).value and ws.cell(r, 1).value not in (None, '') and isinstance(p, (int, float)):
            rows.append(dict(n=ws.cell(r, 1).value, name=str(b).split('\n')[0].strip(), unit=str(ws.cell(r, 4).value).split('/')[0].strip(), qty=q, price=p, total=float(q) * float(p)))
    def q(*k): return next((i['qty'] for i in rows if all(x in i['name'] for x in k)), None)
    def sec_of(n):
        if n.startswith('كابل'): return 2
        if n.startswith('سبلتر') or n.startswith('كلوجر'): return 3
        if n.startswith(('متفرقات', 'صليب', 'مواد تعليق')): return 4
        if n.startswith(('عمليات حفر', 'غرف تفتيش')): return 5
        if n.startswith('مصاريف'): return 6
        if n.startswith('اختبار'): return 7
        return 1
    sec = {i: 0.0 for i in range(1, 8)}
    for i in rows: sec[sec_of(i['name'])] += i['total']
    SUB = round(sum(sec.values()), 2); MARG = round(SUB * STD['contingency'], 2)
    sub_i, marg_i = round(SUB), round(MARG); fin_i = sub_i + marg_i
    si = {k: round(v) for k, v in sec.items()}; si[max(si, key=si.get)] += sub_i - sum(si.values())
    pct = {k: round(si[k] / sub_i * 100, 1) if sub_i else 0 for k in si}
    names = {1: 'المعدات النشطة وتجهيز المقسم (POP)', 2: 'الكابلات الضوئية', 3: 'المقسمات الضوئية والعلب', 4: 'التعليق الهوائي والأعمدة', 5: 'القسم الأرضي', 6: 'أعمال التنفيذ والعمالة', 7: 'الاختبار والقياس والتوثيق'}
    cabs = lambda p: '، '.join(i['name'].split()[2] for i in rows if i['name'].startswith(p))
    exc, man = q('عمليات حفر'), q('غرف تفتيش')
    gv = lambda a, dflt: (bd[a].value if bd is not None and isinstance(bd[a].value, (int, float)) else dflt)
    months = gv('C14', 6)
    return dict(town=town, owner=STD['offer']['second_party'], presenter=STD['offer']['first_party'], date=datetime.date.today().strftime('%d / %m / %Y'), subs=int(gv('C19', 0)), buildings=int(gv('C18', 0)), months=int(months),
        sections=[dict(n=k, name=names[k], val=si[k], pct=pct[k]) for k in range(1, 8)], sub=sub_i, marg=marg_i, fin=fin_i, w_fin=words(fin_i),
        prim=q('سبلتر 1:8 رئيسي'), sec=q('سبلتر 1:8 فرعي'), clo=q('كلوجر'), hasdig=bool(exc), exc=exc, man=man, olt=q('OLT'), board=q('بطاقة GPON'), odf=q('ODF'), patch=q('Patch cord'), cross=q('صليب'),
        cabs_main=cabs('كابل رئيسي'), cabs_br=cabs('كابل فرعي'), items=rows)

# ------------------------------------------------------------------ Word offer (python-docx)
def _design():
    d = STD['offer']['design']; return d['ink'], d['mute'], d['line'], d['accent'], d['soft'], d['gold'], d['soft2'], d['font']
INK = MUTE = LINE = ACC = SOFT = GOLD = SOFT2 = FONT = None   # set from standards when build_docx runs
def build_docx(d, out_path):
    global INK, MUTE, LINE, ACC, SOFT, GOLD, SOFT2, FONT
    INK, MUTE, LINE, ACC, SOFT, GOLD, SOFT2, FONT = _design()
    import docx
    from docx.shared import Pt, Twips, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH as AL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    fmt = lambda n: '{:,}'.format(int(round(n)))
    doc = docx.Document(); sec = doc.sections[0]
    sec.page_width, sec.page_height = Twips(11906), Twips(16838)
    sec.left_margin = sec.right_margin = Twips(850); sec.top_margin = Twips(600); sec.bottom_margin = Twips(700); sec.footer_distance = Twips(300)
    sp = sec._sectPr; bidi = OxmlElement('w:bidi'); sp.append(bidi)
    st = doc.styles['Normal']; st.font.name = FONT; st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn('w:cs'), FONT); st.element.rPr.rFonts.set(qn('w:ascii'), FONT); st.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
    def setrun(r, size=10, bold=False, color=INK):
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color); r.font.name = FONT
        rpr = r._r.get_or_add_rPr(); rpr.append(OxmlElement('w:rtl')); rf = rpr.find(qn('w:rFonts'))
        if rf is None: rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
        for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), FONT)
        if bold: rpr.append(OxmlElement('w:bCs'))
        szcs = OxmlElement('w:szCs'); szcs.set(qn('w:val'), str(int(size * 2))); rpr.append(szcs)
    def para(container, runs, align=None, before=0, after=0, line=None, keep=False):
        p = container.add_paragraph(); pPr = p._p.get_or_add_pPr(); pPr.append(OxmlElement('w:bidi'))
        p.paragraph_format.space_before = Twips(before); p.paragraph_format.space_after = Twips(after)
        if line: p.paragraph_format.line_spacing = Twips(line)
        if keep: p.paragraph_format.keep_with_next = True
        if align == 'c': p.alignment = AL.CENTER
        elif align == 'e': p.alignment = AL.RIGHT if False else AL.LEFT   # in an RTL paragraph LEFT means the end side
        for t, size, bold, color in runs: setrun(p.add_run(t), size, bold, color)
        return p
    def shade(cell, fill):
        tcPr = cell._tc.get_or_add_tcPr(); s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear'); s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), fill); tcPr.append(s)
    def borders(cell, **kw):
        tcPr = cell._tc.get_or_add_tcPr(); b = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            e = OxmlElement('w:' + side); v = kw.get(side)
            if v: e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(v[0])); e.set(qn('w:color'), v[1])
            else: e.set(qn('w:val'), 'nil')
            b.append(e)
        tcPr.append(b)
    def margins(cell, top=70, bottom=70, left=100, right=100):
        tcPr = cell._tc.get_or_add_tcPr(); m = OxmlElement('w:tcMar')
        for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
            e = OxmlElement('w:' + k); e.set(qn('w:w'), str(v)); e.set(qn('w:type'), 'dxa'); m.append(e)
        tcPr.append(m)
    def valign(cell):
        tcPr = cell._tc.get_or_add_tcPr(); v = OxmlElement('w:vAlign'); v.set(qn('w:val'), 'center'); tcPr.append(v)
    def table(widths, nrows=1):
        t = doc.add_table(rows=nrows, cols=len(widths)); t.autofit = False
        tblPr = t._tbl.tblPr; tblPr.append(OxmlElement('w:bidiVisual'))
        lay = OxmlElement('w:tblLayout'); lay.set(qn('w:type'), 'fixed'); tblPr.append(lay)
        tw = tblPr.find(qn('w:tblW'))
        if tw is None: tw = OxmlElement('w:tblW'); tblPr.append(tw)
        tw.set(qn('w:w'), str(sum(widths))); tw.set(qn('w:type'), 'dxa')
        grid = t._tbl.tblGrid
        for i, gc in enumerate(grid.findall(qn('w:gridCol'))): gc.set(qn('w:w'), str(widths[i]))
        for row in t.rows:
            for i, c in enumerate(row.cells): c.width = Twips(widths[i]); borders(c); margins(c); valign(c)
        return t
    def clear(cell):
        for p in list(cell.paragraphs): p._p.getparent().remove(p._p)
    def put(cell, runs, **kw):
        clear(cell) if cell.paragraphs and not cell.paragraphs[0].runs and len(cell.paragraphs) == 1 else None
        return para(cell, runs, **kw)
    def tiny(container, after=0):
        p = para(container, [('', 1, False, INK)], after=after); from docx.enum.text import WD_LINE_SPACING as LS
        p.paragraph_format.line_spacing_rule = LS.EXACTLY; p.paragraph_format.line_spacing = Twips(20); return p
    def gap(h): tiny(doc, h)
    def cant_split(row):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
    def header_row(row):
        trPr = row._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:tblHeader'))
    W = 10206; hw = W // 2
    # letterhead band
    t = table([W], 2); c = t.cell(0, 0); shade(c, ACC); margins(c, 180, 180, 200, 200); clear(c)
    inner = c.add_table(rows=1, cols=2); inner._tbl.tblPr.append(OxmlElement('w:bidiVisual'))
    inner.autofit = False
    for gc in inner._tbl.tblGrid.findall(qn('w:gridCol')): gc.set(qn('w:w'), str((W - 400) // 2))
    for cc in inner.rows[0].cells: cc.width = Twips((W - 400) // 2); borders(cc); margins(cc, 0, 0, 0, 0)
    put(inner.cell(0, 0), [(STD['offer'].get('brand_ar', ''), 19, True, 'FFFFFF')]); para(inner.cell(0, 0), [('%s  ·  FTTH Infrastructure' % STD['offer'].get('brand_en', ''), 7, False, 'AEB8CC')])
    put(inner.cell(0, 1), [('عرض سعر', 19, True, 'FFFFFF')], align='e'); para(inner.cell(0, 1), [('PRICE QUOTATION', 7, False, 'E3B764')], align='e')
    tiny(c)   # table cell must end with a paragraph
    c2 = t.cell(1, 0); shade(c2, GOLD); margins(c2, 0, 0, 0, 0); t.rows[1].height = Twips(70); clear(c2); tiny(c2)
    gap(60)
    t = table([W]); c = t.cell(0, 0); shade(c, SOFT); margins(c, 80, 80, 200, 200); borders(c, right=(24, ACC)); clear(c)
    para(c, [('الموضوع:  ', 10.5, True, ACC), ('تمديد شبكة ألياف ضوئية FTTH هوائية — بلدة %s' % d['town'], 10.5, True, INK)])
    gap(100)
    t = table([hw - 100, 200, hw - 100]); cant_split(t.rows[0])
    def box(cell, title, lines):
        margins(cell, 80, 80, 200, 200); borders(cell, top=(4, LINE), bottom=(4, LINE), left=(4, LINE), right=(4, LINE)); clear(cell); para(cell, [(title, 8, True, GOLD)], after=50)
        for a, b in lines: para(cell, [(a + ':  ', 8.5, False, MUTE), (b, 9.5, True, INK)], after=30)
    box(t.cell(0, 0), 'الأطراف', [('الطرف الأول', d['presenter']), ('الطرف الثاني', d['owner'])])
    box(t.cell(0, 2), 'بيانات العرض', [('تاريخ العرض', d['date']), ('صلاحية العرض', '30 يوماً'), ('مدة التنفيذ', '%d أشهر' % d['months']), ('العملة', 'الدولار الأمريكي (USD)'), ('المرجع', '……………')])
    gap(140)
    HN = [0]
    rowB = dict(bottom=(4, LINE))
    def H(title, en):
        HN[0] += 1; m = title.split(' — ')
        t = table([420, W - 420]); a, b = t.cell(0, 0), t.cell(0, 1)
        shade(a, ACC); margins(a, 40, 40); borders(a, bottom=(12, GOLD)); clear(a); para(a, [(str(HN[0]), 9.5, True, 'FFFFFF')], align='c')
        margins(b, 40, 40); borders(b, bottom=(12, GOLD)); clear(b); para(b, [(m[1] if len(m) > 1 else title, 11.5, True, ACC), ('    ' + en, 8, False, MUTE)])
        gap(80)
    cover = ('منها %s متراً حفراً أرضياً والباقي هوائي' % fmt(d['exc'])) if d['hasdig'] else 'وهي شبكة هوائية بالكامل'
    H('أولاً — الملخص التنفيذي', 'Executive Summary')
    para(doc, [('يغطّي هذا العرض توريد وتنفيذ شبكة ألياف ضوئية FTTH في بلدة %s، %s. تتضمن الشبكة %s نقطة تقسيم رئيسية و%s نقطة تقسيم ثانوية و%s علبة وصل، وتُصمَّم لتغطية %s وخدمة %s مشترك. قيمة العرض الإجمالية %s دولاراً أمريكياً شاملة احتياطي الطوارئ.'
        % (d['town'], cover, fmt(d['prim']), fmt(d['sec']), fmt(d['clo']), (fmt(d['buildings']) + ' منزل') if d['buildings'] else 'البلدة', fmt(d['subs']), fmt(d['fin'])), 10, False, INK)], line=340, after=80)
    gap(80)
    H('ثانياً — ملخص الأقسام', 'Summary by Section')
    cols = [600, 5206, 2400, 2000]
    t = table(cols, 1 + len(d['sections']) + 1); header_row(t.rows[0])
    for i, h in enumerate(('#', 'القسم', 'القيمة (USD)', 'النسبة')):
        cc = t.cell(0, i); shade(cc, ACC); clear(cc); para(cc, [(h, 9, True, 'FFFFFF')], align='c' if i != 1 else None)
    for r, s in enumerate(d['sections'], 1):
        cant_split(t.rows[r])
        for i, v in enumerate((str(s['n']), s['name'], fmt(s['val']) if s['val'] else '—', ('%.1f%%' % s['pct']) if s['val'] else '—')):
            cc = t.cell(r, i); borders(cc, **rowB); clear(cc); para(cc, [(v, 9.5, False, MUTE if i == 0 else INK)], align='c' if i != 1 else None)
    r = len(d['sections']) + 1
    for i, v in enumerate(('', 'المجموع الفرعي للأعمال', fmt(d['sub']), '100%')):
        cc = t.cell(r, i); shade(cc, SOFT); clear(cc); para(cc, [(v, 9.5, True, INK)], align='c' if i != 1 else None)
    doc.add_page_break()
    H('ثالثاً — الخلاصة المالية', 'Financial Summary')
    t = table([6606, 3600], 3)
    for r, (a, b, bold, fill) in enumerate((('المجموع الفرعي للأعمال', fmt(d['sub']) + ' USD', False, None), ('احتياطي الطوارئ (%d%%)' % round(STD['contingency'] * 100), fmt(d['marg']) + ' USD', False, None), ('الإجمالي النهائي للعرض', fmt(d['fin']) + ' USD', True, SOFT2))):
        for i, v in enumerate((a, b)):
            cc = t.cell(r, i); clear(cc)
            if fill: shade(cc, fill)
            else: borders(cc, **rowB)
            para(cc, [(v, 12 if (bold and i) else 10, bold, INK)], align='e' if i else None)
    para(doc, [('المبلغ كتابةً: ', 9.5, True, INK), ('فقط %s دولاراً أمريكياً لا غير.' % d['w_fin'], 9.5, False, INK)], before=100, after=100)
    gap(120)
    H('رابعاً — نطاق العمل المشمول بالعرض', 'Scope of Work')
    scope = ['تجهيز المقسم المركزي: %s جهاز OLT مع %s بطاقات GPON، خزانة شبكة 42U، %s وحدات ODF سعة 48 كور، %s وصلة Patch cord، وحدة طاقة احتياطية وبناء وتجهيز غرفة المقسم.' % (d['olt'], d['board'], d['odf'], d['patch']),
        'Switch Mikrotik CRS326 24S+ 2Q',
        'توريد وتمديد الكابلات الضوئية: رئيسية %s فايبر، وفرعية %s فايبر.' % (d['cabs_main'], d['cabs_br']),
        'توريد وتركيب %s مقسماً رئيسياً و%s مقسماً ثانوياً (1:8)، شاملةً الاحتياط 20%%.' % (fmt(d['prim']), fmt(d['sec'])),
        'توريد وتركيب %s علبة وصل هوائية (كلوجر).' % fmt(d['clo']),
        'أعمال التعليق الهوائي: متفرقات التثبيت و%s صليب تخزين كابلات ومواد التعليق والتثبيت.' % fmt(d['cross']),
        ('القسم الأرضي: %s م حفر وتمديد ودفن وإعادة، و%s غرفة تفتيش مسبقة الصنع.' % (fmt(d['exc']), d['man'])) if d['hasdig'] else 'لا يتضمن العرض أعمالاً أرضية؛ الشبكة هوائية بالكامل.',
        'أعمال التنفيذ والعمالة والإشراف على مدى %s أشهر، شاملةً المحروقات.' % d['months'],
        'المخططات التنفيذية As-Built وملفات QGIS وأعمال اختبار OTDR والقياس.']
    t = table([600, W - 600], len(scope))
    for i, s in enumerate(scope):
        cant_split(t.rows[i]); a, b = t.cell(i, 0), t.cell(i, 1)
        for cc in (a, b): borders(cc, **rowB); margins(cc, 60, 60)
        clear(a); para(a, [('%d.' % (i + 1), 9.5, True, ACC)], align='c'); clear(b); para(b, [(s, 9.5, False, INK)], line=320)
    gap(120)
    H('خامساً — الشروط', 'Terms')
    terms = [('الصلاحية', 'العرض صالح 30 يوماً من تاريخه، وبعدها يخضع لإعادة التسعير.'), ('مدة التنفيذ', '%d أشهر تبدأ من تاريخ استلام الموقع والتراخيص وأذونات المرور، وتُنفَّذ على مراحل.' % d['months']), ('التشغيل', 'يتم التشغيل بشكل مباشر عند انتهاء القطاع المنفذ.')]
    t = table([1800, W - 1800], len(terms))
    for i, (a, b) in enumerate(terms):
        cant_split(t.rows[i]); x, y = t.cell(i, 0), t.cell(i, 1)
        for cc in (x, y): borders(cc, **rowB)
        clear(x); para(x, [(a, 9.5, True, INK)]); clear(y); para(y, [(b, 9.5, False, INK)], line=320)
    doc.add_page_break()
    H('سادساً — ملحق جدول الكميات والأسعار', 'Annex: Bill of Quantities')
    ac = [550, 4506, 1100, 1250, 1100, 1700]; items = d['items']
    t = table(ac, 1 + len(items) + 1); header_row(t.rows[0])
    for i, h in enumerate(('#', 'البند', 'الوحدة', 'الكمية', 'السعر', 'الإجمالي (USD)')):
        cc = t.cell(0, i); shade(cc, ACC); margins(cc, 55, 55); clear(cc); para(cc, [(h, 8, True, 'FFFFFF')], align='c' if i != 1 else None)
    num = lambda v: ('{:,.2f}'.format(v).rstrip('0').rstrip('.')) if isinstance(v, (int, float)) else str(v)
    for r, it in enumerate(items, 1):
        cant_split(t.rows[r])
        for i, v in enumerate((it['n'], it['name'], it['unit'], num(it['qty']), num(it['price']), '{:,.2f}'.format(it['total']))):
            cc = t.cell(r, i); borders(cc, bottom=(2, 'D9D9D9')); margins(cc, 8, 8); clear(cc); para(cc, [(str(v), 8, False, MUTE if i == 0 else INK)], align='c' if i != 1 else None)
    r = len(items) + 1; a = t.cell(r, 0).merge(t.cell(r, 4)); shade(a, SOFT); clear(a); para(a, [('المجموع الفرعي للأعمال', 9, True, INK)])
    z = t.cell(r, 5); shade(z, SOFT); clear(z); para(z, [(fmt(d['sub']), 9, True, INK)], align='c')
    doc.add_page_break()
    H('سابعاً — الاعتماد والتوقيع', 'Acceptance')
    para(doc, [('بتوقيع هذه الصفحة يُقرّ الطرفان بالاطلاع على جدول الكميات المرفق ونطاق العمل والشروط الواردة أعلاه، واعتمادها أساساً للتعاقد.', 9.5, False, INK)], line=330, after=160)
    t = table([hw - 100, 200, hw - 100]); cant_split(t.rows[0])
    for ci, (title, name) in ((0, ('عن الطرف الأول', d['presenter'])), (2, ('عن الطرف الثاني', d['owner']))):
        cc = t.cell(0, ci); borders(cc, top=(4, LINE), bottom=(4, LINE), left=(4, LINE), right=(4, LINE)); margins(cc, 140, 70); clear(cc)
        para(cc, [(title, 9.5, True, INK)], after=40); para(cc, [(name, 9.5, False, INK)], after=90); para(cc, [('الصفة:  ……………………………', 9, False, MUTE)], after=140)
        para(cc, [('التاريخ:  ……… / ……… / ………', 9, False, MUTE)], after=50); para(cc, [('التوقيع والختم:', 9, False, MUTE)], after=900)
    # footer
    fp = sec.footer.paragraphs[0]; fp.text = ''; fp.alignment = AL.CENTER; fp._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
    pb = OxmlElement('w:pBdr'); tp = OxmlElement('w:top'); tp.set(qn('w:val'), 'single'); tp.set(qn('w:sz'), '6'); tp.set(qn('w:space'), '6'); tp.set(qn('w:color'), GOLD); pb.append(tp); fp._p.get_or_add_pPr().append(pb)
    setrun(fp.add_run('%s  ·  عرض سعر FTTH %s  ·  صفحة ' % (STD['offer'].get('brand_ar', ''), d['town'])), 8, False, MUTE)
    r = fp.add_run(); setrun(r, 8, False, MUTE)
    for typ, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        if typ: e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), typ)
        else: e = OxmlElement('w:instrText'); e.text = txt; e.set(qn('xml:space'), 'preserve')
        r._r.append(e)
    # remove first empty paragraph python-docx leaves at the top
    body = doc.element.body
    if body[0].tag == qn('w:p') and not body[0].xpath('.//w:t'): body.remove(body[0])
    doc.core_properties.title = 'عرض سعر FTTH %s' % d['town']; doc.core_properties.author = STD['offer'].get('brand_ar', '')
    doc.save(out_path)
